import re
from pathlib import Path

from docx import Document as WordDocument
from pypdf import PdfReader


class DocumentTextError(ValueError):
    pass


SUPPORTED_ANALYSIS_TYPES = {".pdf", ".docx", ".txt", ".rtf", ".csv"}


def extract_document_text(path: Path) -> str:
    extension = path.suffix.lower()
    if extension not in SUPPORTED_ANALYSIS_TYPES:
        raise DocumentTextError("AI analysis currently supports PDF, DOCX, TXT, RTF, and CSV files.")
    try:
        if extension == ".pdf":
            text = "\n\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
        elif extension == ".docx":
            document = WordDocument(path)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            for table in document.tables:
                for row in table.rows:
                    text += "\n" + " | ".join(cell.text for cell in row.cells)
        else:
            text = path.read_text(encoding="utf-8", errors="replace")
            if extension == ".rtf":
                text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", text)
                text = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", text)
                text = re.sub(r"[{}]", " ", text)
    except Exception as exc:
        raise DocumentTextError(f"Text could not be extracted from this file: {exc}") from exc
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 30:
        raise DocumentTextError("The file does not contain enough extractable text. Scanned PDFs require OCR.")
    return text
